"""
Exporta/sincroniza todos los avisos a un fichero Excel.
"""
import io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from services.avisos import get_all_avisos
from services.comentarios import get_comentarios


_COLOR_MAP = {
    "Acabado":        "C6EFCE",   # verde
    "En proceso":     "BDD7EE",   # azul
    "Falta material": "FFEB9C",   # amarillo/naranja
}
_RED_FILL    = PatternFill("solid", fgColor="FFC7CE")
_ORANGE_FILL = PatternFill("solid", fgColor="FFD966")
_BLUE_FILL   = PatternFill("solid", fgColor="BDD7EE")
_GREEN_FILL  = PatternFill("solid", fgColor="C6EFCE")
_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_BORDER      = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin")
)


def _row_fill(aviso: dict) -> PatternFill:
    if aviso["estado"] == "Acabado":
        return _GREEN_FILL
    from datetime import date
    try:
        fecha = datetime.strptime(aviso["fecha_solicitud"][:10], "%Y-%m-%d").date()
    except Exception:
        return _BLUE_FILL
    dias = (date.today() - fecha).days
    if dias > 90:
        return _RED_FILL
    elif dias > 60:
        return _ORANGE_FILL
    return _BLUE_FILL


def generate_excel() -> bytes:
    """Genera el Excel con todos los avisos y devuelve los bytes."""
    avisos = get_all_avisos()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Avisos"

    headers = [
        "Aviso", "Fecha de solicitud", "Generador OT", "Generador Aviso",
        "E.S.M.", "Descripción de la OT",
        "Estado", "Coordinador asignado", "Enlace Drive",
        "Material necesario", "Fecha cierre", "Última actualización",
        "Comentarios",
    ]

    # Cabecera
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill    = _HEADER_FILL
        cell.font    = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border  = _BORDER

    ws.row_dimensions[1].height = 30

    # Filas
    for row_idx, av in enumerate(avisos, 2):
        comentarios = get_comentarios(av["id"])
        coment_text = " | ".join(
            f"[{c['fecha_comentario'][:16]}] {c['autor_nombre']}: {c['texto']}"
            for c in comentarios
        )

        values = [
            av["num_aviso"],
            av["fecha_solicitud"],
            av["generador_ot"],
            av["generador_aviso"],
            av["esm"],
            av["descripcion"],
            av["estado"],
            av["coordinador_nombre"] or "",
            av["enlace_drive"] or "",
            av["material_necesario"] or "",
            av["fecha_cierre"] or "",
            av["updated_at"][:16] if av.get("updated_at") else "",
            coment_text,
        ]

        fill = _row_fill(av)
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.fill      = fill
            cell.border    = _BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 6 or col == 13))

    # Anchos de columna
    col_widths = [10, 16, 22, 22, 55, 60, 16, 22, 35, 35, 14, 18, 80]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
